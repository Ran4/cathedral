# /// script
# requires-python = ">=3.11"
# dependencies = ["python-docx==1.2.0", "matplotlib==3.10.8", "pillow==12.1.1", "pymupdf==1.27.1"]
# ///
"""Build An Alibi in Stone using python-docx; convert, validate, optionally open.

    uv run --cache-dir /tmp/cathedral-gdd-uv generate_gdd.py --open

Source text lives in manuscript.md. Original screenshots/art are copied unchanged.
Figures are code-authored diagrams, not AI-generated raster artwork.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile

sys.dont_write_bytecode = True

os.environ.setdefault("MPLCONFIGDIR", "/tmp/cathedral-gdd-matplotlib")
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch, Circle, Rectangle
import pymupdf
from PIL import Image, ImageDraw
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
REPO = HERE.parent.parent
FIGURES = HERE / "figures"
PREVIEWS = HERE / "previews"
STEM = "an_alibi_in_stone_gdd"
EXPECTED_PAGES = 35
DIAGRAM_COUNT = 9
PAPER = "#FBF8F1"
INK = "#263D42"
GREEN = "#426B60"
GOLD = "#B3873F"
RED = "#AA604A"
PALE = "#E8EEE8"
SAND = "#F0E4CB"
MUTED = "#626E70"
LINE = "#C7D1CB"

plt.rcParams.update({
    "font.family": "DejaVu Sans", "font.size": 10,
    "text.color": INK, "axes.labelcolor": INK,
    "figure.facecolor": PAPER, "axes.facecolor": PAPER,
    "svg.fonttype": "none",
})


def canvas(height=4):
    fig = plt.figure(figsize=(10, height))
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set(xlim=(0, 10), ylim=(0, height))
    ax.axis("off")
    return fig, ax


def label(ax, x, y, value, size=11, color=INK, weight="normal", ha="center", va="center"):
    return ax.text(x, y, value, fontsize=size, color=color, weight=weight,
                   ha=ha, va=va, linespacing=1.4)


def box(ax, x, y, w, h, title, body="", color=PALE, edge=GREEN, size=11):
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.025,rounding_size=0.07",
                              facecolor=color, edgecolor=edge, linewidth=1.2))
    label(ax, x+w/2, y+h-(0.28 if body else h/2), title, size=size, weight="bold")
    if body:
        label(ax, x+w/2, y+(h-0.45)/2, body, size=size-1, color=MUTED)


def arrow(ax, start, end, color=GREEN, rad=0, style="-"):
    ax.add_patch(FancyArrowPatch(start, end, arrowstyle="-|>", mutation_scale=14,
                                color=color, lw=1.6, linestyle=style,
                                connectionstyle=f"arc3,rad={rad}"))


def save_figure(fig, name):
    for extension in ("png", "svg"):
        fig.savefig(FIGURES / f"{name}.{extension}", dpi=220, facecolor=PAPER,
                    bbox_inches=None, pad_inches=0)
    plt.close(fig)


def make_figures():
    from quest_figures import make_figures as render_quest_figures
    FIGURES.mkdir(exist_ok=True)
    render_quest_figures(canvas, box, label, arrow, save_figure)


def shade(element, fill):
    shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),fill.lstrip("#")); element.append(shd)


def inline(paragraph, value, size=None):
    for part in re.split(r"(\*\*.*?\*\*)",value):
        run=paragraph.add_run(part[2:-2] if part.startswith("**") else part)
        run.bold=part.startswith("**")
        if size: run.font.size=Pt(size)


def field(p, code):
    r=p.add_run(); f=OxmlElement("w:fldSimple"); f.set(qn("w:instr"),code); r._r.addnext(f)


def add_picture(doc, name, caption, max_height):
    path=FIGURES/name
    with Image.open(path) as im: aspect=im.width/im.height
    width=min(6.9,max_height*aspect)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(3)
    p.paragraph_format.keep_with_next=True
    shape=p.add_run().add_picture(str(path),width=Inches(width))
    shape._inline.docPr.set("descr",caption)
    cap=doc.add_paragraph(caption,"Caption")
    cap.paragraph_format.space_after=Pt(9)


def add_table(doc, lines, page_number):
    rows=[[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
    rows=[r for r in rows if not all(re.fullmatch(r"[:\- ]+",c) for c in r)]
    table=doc.add_table(rows=1,cols=len(rows[0]))
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False
    count=len(rows[0]); total=17.6
    fractions=([.40,.60] if count==2 else [.26,.37,.37] if count==3 else [.18,.24,.32,.26])
    if rows[0][0] == "ID":
        fractions = [.08, .39, .53] if count == 3 else [.08, .35, .31, .26]
    for col,frac in zip(table.columns,fractions): col.width=Cm(total*frac)
    for i,row in enumerate(rows):
        cells=table.rows[0].cells if i==0 else table.add_row().cells
        for j,(cell,content) in enumerate(zip(cells,row)):
            cell.width=Cm(total*fractions[j])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            shade(tcpr,INK if i==0 else (PALE if i%2 else PAPER))
            margins=OxmlElement("w:tcMar")
            for side,v in (("top",80),("bottom",80),("left",95),("right",95)):
                node=OxmlElement("w:"+side); node.set(qn("w:w"),str(v)); node.set(qn("w:type"),"dxa"); margins.append(node)
            tcpr.append(margins)
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=1.06
            inline(p,content,size=9.2 if page_number not in (2,26) else 9.0)
            for run in p.runs:
                run.font.name="Lato"
                if i==0: run.bold=True; run.font.color.rgb=RGBColor.from_string("FFFFFF")
        trpr=table.rows[i]._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if i==0: trpr.append(OxmlElement("w:tblHeader"))
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(1)
    p.paragraph_format.space_before=Pt(0); p.paragraph_format.line_spacing=Pt(2)
    p.add_run().font.size=Pt(2)


def build_docx(chapters):
    doc=Document()
    sec=doc.sections[0]
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.75); sec.bottom_margin=Cm(1.65)
    sec.left_margin=Cm(1.7); sec.right_margin=Cm(1.7)
    sec.header_distance=Cm(.7); sec.footer_distance=Cm(.7)
    sec.different_first_page_header_footer=True
    styles=doc.styles
    for name in ("Normal","Body Text"):
        st=styles[name]; st.font.name="Lato"; st.font.size=Pt(10.5)
        st.font.color.rgb=RGBColor.from_string(INK[1:])
        st.paragraph_format.line_spacing=1.13
        st.paragraph_format.space_after=Pt(7)
        st.paragraph_format.widow_control=True
    for name,size,color in (("Title",42,INK),("Heading 1",25,INK),("Heading 2",13,GREEN)):
        st=styles[name]; st.font.name="EB Garamond" if name!="Heading 2" else "Lato"
        st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(color[1:])
        st.font.bold=name=="Heading 2"
        st.paragraph_format.space_before=Pt(7 if name=="Heading 2" else 0)
        st.paragraph_format.space_after=Pt(7 if name=="Heading 2" else 12)
        st.paragraph_format.keep_with_next=True
        fonts=st.element.find(qn("w:rPr")).find(qn("w:rFonts"))
        for attr in list(fonts.attrib):
            if "theme" in attr.lower(): del fonts.attrib[attr]
    # The packaged Word template supplies a blue title rule; use our palette.
    for border in styles["Title"].element.iter(qn("w:bottom")):
        border.set(qn("w:color"),GOLD[1:])
        for attr in list(border.attrib):
            if "theme" in attr.lower(): del border.attrib[attr]
    styles["Caption"].font.name="Lato"; styles["Caption"].font.size=Pt(8.2)
    styles["Caption"].font.color.rgb=RGBColor.from_string(MUTED[1:])
    styles["Caption"].font.italic=False
    styles["Caption"].paragraph_format.line_spacing=1.08
    styles["List Bullet"].font.name="Lato"; styles["List Bullet"].font.size=Pt(10.5)
    styles["List Bullet"].paragraph_format.space_after=Pt(4)
    styles["List Bullet"].paragraph_format.line_spacing=1.10
    header=sec.header.paragraphs[0]
    header.text="OMBREVAL     /     AN ALIBI IN STONE"
    header.runs[0].font.name="Lato"; header.runs[0].font.size=Pt(8)
    header.runs[0].font.color.rgb=RGBColor.from_string(GREEN[1:])
    borders=OxmlElement("w:pBdr"); bottom=OxmlElement("w:bottom")
    for k,v in (("val","single"),("sz","6"),("color",LINE[1:]),("space","7")): bottom.set(qn("w:"+k),v)
    borders.append(bottom); header._p.get_or_add_pPr().append(borders)
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=footer.add_run("DESIGN PROPOSAL  ·  04 SEPTEMBER 2026       ")
    run.font.name="Lato"; run.font.size=Pt(7.5); run.font.color.rgb=RGBColor.from_string(MUTED[1:])
    field(footer,"PAGE"); footer.add_run(" / "); field(footer,"NUMPAGES")
    props=doc.core_properties
    props.title="An Alibi in Stone — A complete investigation quest for Ombreval"
    props.subject="A fixed crime, a false continuous alibi, a walkable reconstruction, and independently supported findings"
    props.author="Codex"
    props.keywords="Ombreval; Cathedral; game design; investigation; alibi; quest"
    props.comments="Repository-grounded proposal. No new gameplay is implemented by this document."

    p=doc.add_paragraph("THE CATHEDRAL-CITY OF IMPOSSIBLE LIGHT")
    p.paragraph_format.space_before=Pt(14)
    p.runs[0].font.size=Pt(10); p.runs[0].font.bold=True
    p.runs[0].font.color.rgb=RGBColor.from_string(GREEN[1:])
    doc.add_paragraph("AN ALIBI\nIN STONE", "Title")
    p=doc.add_paragraph("Walk the claim. Find the missing route.\nProve who used it.")
    p.paragraph_format.space_after=Pt(19)
    for r in p.runs: r.font.name="EB Garamond"; r.font.size=Pt(19)
    add_picture(doc,"reference_image.png","The project's existing architectural inspiration. The crime, rooms and routes in this document are proposals; this image is not their surveyed map.",5.10)
    p=doc.add_paragraph("A complete investigation quest for Ombreval")
    p.runs[0].bold=True; p.paragraph_format.space_before=Pt(7)
    doc.add_paragraph("04 September 2026  ·  Quest GDD  ·  35 pages")
    p=doc.add_paragraph("An assault, a stolen packet and two genuine sightings. Reconstruct a crime through the movements, work and hidden connections of a living city.")
    p.paragraph_format.space_after=Pt(0)

    for number,chapter in enumerate(chapters,2):
        lines=chapter.strip().splitlines(); title=lines[0][2:]
        p=doc.add_paragraph(title,"Heading 1"); p.paragraph_format.page_break_before=True
        start=OxmlElement("w:bookmarkStart"); start.set(qn("w:id"),str(number)); start.set(qn("w:name"),f"page_{number}")
        end=OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"),str(number)); p._p.insert(0,start); p._p.append(end)
        i=1
        while i<len(lines):
            line=lines[i].strip()
            if not line: i+=1; continue
            if line.startswith("|"):
                table_lines=[]
                while i<len(lines) and lines[i].startswith("|"): table_lines.append(lines[i]); i+=1
                add_table(doc,table_lines,number); continue
            if line.startswith("{{figure:"):
                name,caption,height=line[len("{{figure:"):-2].split("|")
                add_picture(doc,name,caption,float(height)); i+=1; continue
            if line.startswith("## "):
                doc.add_paragraph(line[3:],"Heading 2"); i+=1; continue
            if line.startswith("- "):
                inline(doc.add_paragraph(style="List Bullet"),line[2:]); i+=1; continue
            if line.startswith("> "):
                p=doc.add_paragraph(); inline(p,line[2:])
                p.paragraph_format.left_indent=Cm(.35); p.paragraph_format.right_indent=Cm(.3)
                p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(10)
                shade(p._p.get_or_add_pPr(),SAND); i+=1; continue
            p=doc.add_paragraph()
            # Insert break opportunities only inside long source-path tokens.
            if number>=34:
                line=re.sub(r"\S+/\S+",lambda m:m[0].replace("/","/\u200b").replace("_","_\u200b"),line)
                p.paragraph_format.line_spacing=1.1
            inline(p,line,size=9.7 if number>=34 else None)
            i+=1
    path=HERE/f"{STEM}.docx"; doc.save(path)
    return path


SOURCE_PATHS = ['AGENTS.md', 'features/AGENTS.md', 'crates/cathedral-sim/AGENTS.md', 'config.ron', 'crates/cathedral-sim/src/lib.rs', 'crates/cathedral-sim/src/engine.rs', 'crates/cathedral-sim/src/actions.rs', 'crates/cathedral-sim/src/world.rs', 'crates/cathedral-sim/src/round.rs', 'crates/cathedral-sim/src/clock.rs', 'crates/cathedral-sim/src/night.rs', 'crates/cathedral-sim/src/traits.rs', 'crates/cathedral-sim/src/snapshot.rs', 'crates/cathedral-sim/src/event.rs', 'crates/cathedral-sim/src/notices.rs', 'crates/cathedral-sim/src/custody.rs', 'crates/cathedral-sim/src/marks.rs', 'src/controller.rs', 'src/map.rs', 'src/smart_actors/local_engine.rs', 'src/smart_actors/interaction.rs', 'src/smart_actors/chat.rs', 'src/smart_actors/inventory_ui.rs', 'lore/characters/scribe_and_clerk/fc9rn_corin_copp.json', 'lore/characters/money_dealer/fl5cp_lise_copp.json', 'lore/characters/court_officer/fo6gl_odo_trask.json', 'lore/characters/cargo_worker/fw7ub_warin_underbridge.json', 'lore/characters/revenue_worker/fa8tn_averil_tarn.json', 'lore/characters/scribe_and_clerk/fg4br_gile_of_brede.json', 'lore/characters/bailiff_and_gaoler/p00a3_segwin_mott.json', 'lore/places/00_city_plan.md', 'lore/places/02_canonical_gazetteer.md', 'assets/world/areas.json', 'lore/the_dry_boatmen.md', 'lore/second_sun/00_canon.md', 'features/knowledge_and_rumor/README.md', 'features/knowledge_and_rumor/plan/M1.md', 'features/knowledge_and_rumor/plan/M5.md', 'features/knowledge_and_rumor/m0_evidence/NOTES.md', 'features/keys_and_locked_places.md', 'features/implemented/law_and_order.md', 'features/quest_the_bale_that_gained_forty_pounds/README.md', 'features/quest_ring_a_dead_womans_name_at_marenstide/README.md', 'features/quest_secure_votes_for_a_drainage_funding_plan_before_the_rain/README.md', 'features/systemic_quest_suggestions.md']
IMAGE_SOURCES={"reference_image.png":"docs/reference_image.png"}

def provenance():
    entries=[]
    for relative in SOURCE_PATHS:
        path=REPO/relative
        entries.append({"path":relative,"sha256":hashlib.sha256(path.read_bytes()).hexdigest(),
                        "scope":"status, relevant sections or representative code inspected; not a full-file audit"})
    images=[]
    for target,relative in IMAGE_SOURCES.items():
        origin=REPO/relative; local=FIGURES/target
        if origin.exists(): shutil.copyfile(origin,local)
        elif not local.exists(): raise FileNotFoundError(origin)
        images.append({"file":f"figures/{target}","source":relative,"altered":False,
                       "sha256":hashlib.sha256(local.read_bytes()).hexdigest()})
    result={"inspection_date":"2026-09-04","head":subprocess.check_output(["git","rev-parse","HEAD"],cwd=REPO,text=True).strip(),
            "scope":"Local working tree, including uncommitted design work; no fresh gameplay run", "sources":entries,"images":images,
            "diagrams":"Nine original investigation diagrams; PNG and SVG; proposed topology and timing, not game measurements"}
    result["generation_inputs"] = {name: hashlib.sha256((HERE/name).read_bytes()).hexdigest() for name in ("manuscript.md", "generate_gdd.py", "quest_figures.py", "quest_model.json", "validate_quest.py")}
    (HERE/"source_manifest.json").write_text(json.dumps(result,indent=2)+"\n")


def normalize(value):
    return re.sub(r"\s+"," ",value.replace("\u200b","").replace("\ufb01","fi").replace("\ufb02","fl")).strip()


def validate_and_preview(pdf_path,chapters):
    PREVIEWS.mkdir(exist_ok=True)
    for old in PREVIEWS.iterdir():
        if re.fullmatch(r"(?:page|contact)_\d+\.png", old.name):
            old.unlink()
    pdf=pymupdf.open(pdf_path)
    titles=["AN ALIBI"]+[c.strip().splitlines()[0][2:] for c in chapters]
    pages=[]; issues=[]; thumbs=[]
    for index,page in enumerate(pdf):
        text=page.get_text()
        expected=titles[index] if index<len(titles) else "UNEXPECTED OVERFLOW PAGE"
        match=normalize(expected) in normalize(text)
        if not match: issues.append(f"Page {index+1}: expected title {expected!r} missing")
        outliers=[]
        for block in page.get_text("dict")["blocks"]:
            if block["type"]!=0: continue
            for line in block["lines"]:
                for span in line["spans"]:
                    x0,y0,x1,y1=span["bbox"]
                    if x0<25 or x1>page.rect.width-25 or y0<12 or y1>page.rect.height-12:
                        outliers.append(span["text"])
        if outliers: issues.append(f"Page {index+1}: text outside safe page bounds: {outliers}")
        pages.append({"page":index+1,"expected_heading":expected,"heading_found":match,"words":len(text.split()),"out_of_bounds":outliers})
        pix=page.get_pixmap(matrix=pymupdf.Matrix(1.1,1.1),alpha=False)
        path=PREVIEWS/f"page_{index+1:02}.png"; pix.save(path)
        im=Image.open(path).convert("RGB"); im.thumbnail((210,300))
        thumb=Image.new("RGB",(230,330),"#E3E8E3"); thumb.paste(im,((230-im.width)//2,10))
        ImageDraw.Draw(thumb).text((12,312),str(index+1),fill=INK)
        thumbs.append(thumb)
    for start in range(0,len(thumbs),15):
        sheet=Image.new("RGB",(230*5,330*3),"#E3E8E3")
        for j,thumb in enumerate(thumbs[start:start+15]): sheet.paste(thumb,((j%5)*230,(j//5)*330))
        sheet.save(PREVIEWS/f"contact_{start//15+1:02}.png")
    if len(pdf)!=EXPECTED_PAGES: issues.append(f"Expected {EXPECTED_PAGES} pages, rendered {len(pdf)}")
    result={"page_count":len(pdf),"expected_pages":EXPECTED_PAGES,"diagram_count":DIAGRAM_COUNT,"pages":pages,"issues":issues,
            "checks":"Page count, each expected heading on its page, text bounds, preview generation. Human visual review is separate.",
            "pdf_sha256":hashlib.sha256(pdf_path.read_bytes()).hexdigest()}
    (HERE/"validation.json").write_text(json.dumps(result,indent=2)+"\n")
    print(json.dumps({"pdf":str(pdf_path),"pages":len(pdf),"words":sum(p["words"] for p in pages),"issues":issues},indent=2))
    return not issues


def main():
    parser=argparse.ArgumentParser(); parser.add_argument("--open",action="store_true")
    args=parser.parse_args()
    chapters=(HERE/"manuscript.md").read_text().split("<!-- page -->")
    assert len(chapters)==EXPECTED_PAGES-1, f"Expected {EXPECTED_PAGES-1} content pages plus cover, got {len(chapters)}"
    from validate_quest import validate
    if not validate(): raise SystemExit("Written quest model failed validation")
    make_figures(); provenance(); docx=build_docx(chapters)
    with tempfile.TemporaryDirectory(prefix="cathedral-gdd-soffice-") as profile:
        subprocess.run(["soffice",f"-env:UserInstallation={Path(profile).as_uri()}","--headless","--convert-to","pdf","--outdir",str(HERE),str(docx)],check=True,timeout=120,
                       env={**os.environ,"GSETTINGS_BACKEND":"memory","NO_AT_BRIDGE":"1"})
    pdf=HERE/f"{STEM}.pdf"
    # Opening is explicitly requested by project instructions after first generation.
    if args.open:
        with open(Path(tempfile.gettempdir())/"cathedral-gdd-pdf-open.log","w") as log:
            subprocess.Popen(["xdg-open",str(pdf)],stdout=log,stderr=log,start_new_session=True,
                             env={**os.environ,"GSETTINGS_BACKEND":"memory","NO_AT_BRIDGE":"1"})
    valid=validate_and_preview(pdf,chapters)
    if not valid: raise SystemExit("Layout validation failed; inspect validation.json and previews, then adjust.")


if __name__=="__main__": main()
